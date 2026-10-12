import hashlib, re, tempfile, uuid
from pathlib import Path
from .models import Chunk

def split_text(text,size=850,overlap=120):
    text=re.sub(r'\s+',' ',text).strip()
    if not text: return []
    out=[]; start=0
    while start<len(text):
        end=min(len(text),start+size)
        if end<len(text):
            boundary=max(text.rfind('. ',start,end),text.rfind(' ',start,end)); end=boundary+1 if boundary>start+150 else end
        out.append(text[start:end].strip())
        if end >= len(text): break
        start=max(end-overlap,start+1)
    return out


_CUE=re.compile(r'(?:(\d{1,2}):)?(\d{1,2}):(\d{2})[.,]\d{1,3}\s*-->')
_LEAD=re.compile(r'^\[?(?:(\d{1,2}):)?(\d{1,2}):(\d{2})\]?\s+(.+)$')

def _secs(h,m,s): return int(h or 0)*3600+int(m)*60+int(s)

def parse_transcript(raw):
    """Return [(start_seconds|None, text)] using real VTT/SRT cues or leading timestamps; never guesses."""
    segs=[]; state={'start':None,'buf':[]}
    def flush():
        text=' '.join(state['buf']).strip()
        if text: segs.append((state['start'],text))
        state['buf']=[]
    for line in raw.splitlines():
        line=line.strip()
        if not line or line.upper().startswith(('WEBVTT','NOTE')) or line.isdigit(): continue
        m=_CUE.search(line)
        if m: flush(); state['start']=_secs(*m.groups()); continue
        m=_LEAD.match(line)
        if m: flush(); state['start']=_secs(*m.groups()[:3]); state['buf'].append(m.group(4)); continue
        state['buf'].append(re.sub(r'<[^>]+>','',line))
    flush(); return segs

def transcript_parts(raw,window=700):
    """Group consecutive segments into windows that keep the start time of their first segment."""
    windows=[]
    for start,text in parse_transcript(raw):
        if windows and (windows[-1][0] is None)==(start is None) and len(windows[-1][1])+len(text)<window: windows[-1][1]+=' '+text
        else: windows.append([start,text])
    return [(None,text,(start,None) if start is not None else None) for start,text in windows]

def make_chunks(name,content_type,parts):
    doc=hashlib.sha1(name.encode()).hexdigest()[:12]; chunks=[]
    for location,text,timing in parts:
        for piece in split_text(text):
            chunks.append(Chunk(id=str(uuid.uuid4()),document_id=doc,document_name=name,content_type=content_type,text=piece,page_number=location if content_type=='pdf' else None,slide_number=location if content_type=='pptx' else None,start_seconds=timing[0] if timing else None,end_seconds=timing[1] if timing else None))
    return chunks

def ingest_upload(upload):
    name=upload.name; ext=Path(name).suffix.lower(); data=upload.getvalue()
    if ext=='.pdf':
        import fitz
        doc=fitz.open(stream=data,filetype='pdf'); parts=[(i+1,p.get_text('text'),None) for i,p in enumerate(doc)]
        if not any(t.strip() for _,t,_ in parts): raise ValueError('No selectable text was found. This appears scanned; OCR is required before indexing.')
        return make_chunks(name,'pdf',parts)
    if ext=='.pptx':
        from pptx import Presentation
        with tempfile.NamedTemporaryFile(suffix='.pptx',delete=False) as f: f.write(data); path=f.name
        prs=Presentation(path); return make_chunks(name,'pptx',[(i+1,' '.join(s.text for s in slide.shapes if hasattr(s,'text')),None) for i,slide in enumerate(prs.slides)])
    if ext=='.docx':
        from docx import Document
        with tempfile.NamedTemporaryFile(suffix='.docx',delete=False) as f: f.write(data); path=f.name
        return make_chunks(name,'docx',[(1,' '.join(p.text for p in Document(path).paragraphs),None)])
    if ext in ('.txt','.vtt','.srt'):
        return make_chunks(name,'transcript',transcript_parts(data.decode('utf-8',errors='replace')))
    if ext in ('.mp4','.mov','.mkv','.webm'): raise ValueError('Video needs FFmpeg plus a configured transcription service. Upload its TXT/VTT transcript for the offline workflow.')
    raise ValueError('Unsupported file. Upload PDF, PPTX, DOCX, TXT, VTT, or SRT.')
